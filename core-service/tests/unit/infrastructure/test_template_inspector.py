from io import BytesIO

import pytest
from docx import Document

from src.domain.exceptions.application_templates import InvalidApplicationTemplate
from src.infrastructure.docx.template_inspector import DocxTemplateInspector


def create_docx() -> bytes:
    document = Document()
    document.add_paragraph("Dziekan Wydziału")
    document.add_paragraph("")
    document.add_paragraph("{{ student_name }}")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Numer albumu"
    table.rows[0].cells[1].text = "{{ student_id }}"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_get_text():
    text = await DocxTemplateInspector().get_text(create_docx())

    assert text == "Dziekan Wydziału\n{{ student_name }}\nNumer albumu\n{{ student_id }}"


async def test_get_text_invalid_template():
    with pytest.raises(InvalidApplicationTemplate):
        await DocxTemplateInspector().get_text(b"not a docx")
