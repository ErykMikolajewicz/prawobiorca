import math
from io import BytesIO
from unittest.mock import ANY
from uuid import UUID, uuid4

import pytest
from docx import Document
from fastapi import status
from sqlalchemy import delete, insert, select

from src.app.dtos.application_templates import ApplicationTemplateUploadTarget
from src.domain.value_objects.application_templates import (
    ApplicationFieldType,
    ApplicationTemplateDescriptionStatus,
    ApplicationTemplateField,
    ApplicationTemplateStatus,
)
from src.framework.dependencies.ai_services import get_texts_embedder
from src.infrastructure.relational_db.schemas.application_templates import application_templates_table
from src.infrastructure.relational_db.schemas.organizations import organizations_table, suborganizations_table
from src.main import prawobiorca
from src.shared.consts import ACCESS_COOKIE_NAME, VECTOR_LENGTH
from tests.consts import ACCESS_TOKEN

FIELDS = [
    ApplicationTemplateField(name="department", label="department"),
    ApplicationTemplateField(name="student_id", label="student_id"),
]

FIELDS_JSON = [
    {
        "name": "department",
        "label": "department",
        "fieldType": "TEXT",
        "required": True,
        "passToAi": True,
        "defaultValue": None,
        "pattern": None,
        "options": [],
    },
    {
        "name": "student_id",
        "label": "student_id",
        "fieldType": "TEXT",
        "required": True,
        "passToAi": True,
        "defaultValue": None,
        "pattern": None,
        "options": [],
    },
]


def create_docx_template() -> bytes:
    document = Document()
    document.add_paragraph("{{ student_id }}, {{ department }}, {{ current_date }}")
    document.add_paragraph("{%p for paragraph in paragraphs %}")
    document.add_paragraph("{{ paragraph }}")
    document.add_paragraph("{%p endfor %}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def insert_template(
    session_maker,
    name: str = "Wniosek",
    template_status: ApplicationTemplateStatus = ApplicationTemplateStatus.DRAFT,
    fields: list[ApplicationTemplateField] | None = None,
    instructions: str | None = None,
    organization_id: UUID | None = None,
    suborganization_id: UUID | None = None,
    description: str | None = None,
    description_vector: list[float] | None = None,
):
    async with session_maker.begin() as session:
        statement = (
            insert(application_templates_table)
            .values(
                name=name,
                status=template_status,
                fields=fields or [],
                instructions=instructions,
                organization_id=organization_id,
                suborganization_id=suborganization_id,
                description=description,
                description_vector=description_vector,
            )
            .returning(application_templates_table.c.id)
        )
        return await session.scalar(statement)


async def delete_templates(session_maker, *template_ids):
    async with session_maker.begin() as session:
        await session.execute(
            delete(application_templates_table).where(application_templates_table.c.id.in_(template_ids))
        )


async def insert_organization(session_maker, short_name):
    async with session_maker.begin() as session:
        return await session.scalar(
            insert(organizations_table)
            .values(name=f"Uczelnia {short_name}", short_name=short_name)
            .returning(organizations_table.c.id)
        )


async def insert_suborganization(session_maker, organization_id, name):
    async with session_maker.begin() as session:
        return await session.scalar(
            insert(suborganizations_table)
            .values(organization_id=organization_id, name=name)
            .returning(suborganizations_table.c.id)
        )


async def delete_organizations(session_maker, *organization_ids):
    async with session_maker.begin() as session:
        await session.execute(delete(organizations_table).where(organizations_table.c.id.in_(organization_ids)))


async def get_template_organization(session_maker, template_id):
    async with session_maker() as session:
        result = await session.execute(
            select(
                application_templates_table.c.organization_id, application_templates_table.c.suborganization_id
            ).where(application_templates_table.c.id == template_id)
        )
        return tuple(result.one())


class StubTextsEmbedder:
    @staticmethod
    async def embed_queries(queries):
        return [QUERY_VECTOR]


QUERY_VECTOR = [1.0] * VECTOR_LENGTH
HALF_RELATED_VECTOR = [1.0] * (VECTOR_LENGTH // 2) + [0.0] * (VECTOR_LENGTH // 2)
UNRELATED_VECTOR = [1.0, *[0.0] * (VECTOR_LENGTH - 1)]


async def get_template_description(session_maker, template_id):
    async with session_maker() as session:
        result = await session.execute(
            select(
                application_templates_table.c.description,
                application_templates_table.c.description_vector,
                application_templates_table.c.description_status,
            ).where(application_templates_table.c.id == template_id)
        )
        return result.one()


async def test_get_published_application_templates(
    client, override_session_maker, session_maker, override_authorize_normal_user
):
    published_id = await insert_template(
        session_maker, "Published", ApplicationTemplateStatus.PUBLISHED, FIELDS, "Instrukcje"
    )
    draft_id = await insert_template(session_maker, "Draft", ApplicationTemplateStatus.DRAFT, FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get("/api/application-templates/published")

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == [
            {
                "id": str(published_id),
                "name": "Published",
                "description": None,
                "fields": [{key: value for key, value in field.items() if key != "passToAi"} for field in FIELDS_JSON],
            }
        ]
    finally:
        await delete_templates(session_maker, published_id, draft_id)


async def test_get_application_templates_as_admin(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS, instructions="Instrukcje")

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get("/api/application-templates")

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == [
            {
                "id": str(template_id),
                "createDate": ANY,
                "name": "Wniosek",
                "status": "DRAFT",
                "fields": FIELDS_JSON,
                "instructions": "Instrukcje",
                "description": None,
                "descriptionStatus": "NOT_REQUESTED",
                "organizationId": None,
                "suborganizationId": None,
            }
        ]
    finally:
        await delete_templates(session_maker, template_id)


async def test_get_application_template_as_admin(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS, instructions="Instrukcje")

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get(f"/api/application-templates/{template_id}")

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == {
            "id": str(template_id),
            "createDate": ANY,
            "name": "Wniosek",
            "status": "DRAFT",
            "fields": FIELDS_JSON,
            "instructions": "Instrukcje",
            "description": None,
            "descriptionStatus": "NOT_REQUESTED",
            "organizationId": None,
            "suborganizationId": None,
        }
    finally:
        await delete_templates(session_maker, template_id)

    response = await client.get(f"/api/application-templates/{template_id}")

    assert response.status_code == status.HTTP_404_NOT_FOUND


async def test_get_application_templates_as_normal_user(client, override_session_maker, override_authorize_normal_user):
    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    response = await client.get("/api/application-templates")

    assert response.status_code == status.HTTP_403_FORBIDDEN


async def test_add_application_template(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_templates_storage,
    mock_application_templates_storage,
):
    mock_application_templates_storage.get_upload_target.side_effect = lambda id_: ApplicationTemplateUploadTarget(
        id=id_, url="http://storage.local/bucket", fields={"key": str(id_)}
    )

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    response = await client.post("/api/application-templates", json={"name": "Wniosek"})

    assert response.status_code == status.HTTP_201_CREATED
    template_id = UUID(response.json()["id"])

    try:
        async with session_maker() as session:
            result = await session.execute(
                select(application_templates_table).where(application_templates_table.c.id == template_id)
            )
        template = result.one()

        assert template.name == "Wniosek"
        assert template.status == ApplicationTemplateStatus.NOT_UPLOADED
        assert template.fields == []
        assert template.instructions
    finally:
        await delete_templates(session_maker, template_id)


async def test_confirm_application_template_upload(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_templates_storage,
    mock_application_templates_storage,
):
    template_id = await insert_template(session_maker, template_status=ApplicationTemplateStatus.NOT_UPLOADED)
    mock_application_templates_storage.get_template.return_value = create_docx_template()

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/confirm-upload")

        assert response.status_code == status.HTTP_200_OK
        assert response.json()["status"] == "DRAFT"
        assert response.json()["fields"] == FIELDS_JSON
        mock_application_templates_storage.get_template.assert_awaited_once_with(template_id)
    finally:
        await delete_templates(session_maker, template_id)


async def test_confirm_application_template_upload_invalid_file(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_templates_storage,
    mock_application_templates_storage,
):
    template_id = await insert_template(session_maker, template_status=ApplicationTemplateStatus.NOT_UPLOADED)
    mock_application_templates_storage.get_template.return_value = b"not a docx"

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/confirm-upload")

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    finally:
        await delete_templates(session_maker, template_id)


async def test_get_application_template_download_url(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_templates_storage,
    mock_application_templates_storage,
):
    template_id = await insert_template(session_maker, fields=FIELDS)
    not_uploaded_template_id = await insert_template(
        session_maker, template_status=ApplicationTemplateStatus.NOT_UPLOADED
    )
    mock_application_templates_storage.get_download_url.return_value = "http://storage.local/bucket/template"

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get(f"/api/application-templates/{template_id}/download-url")

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == "http://storage.local/bucket/template"
        mock_application_templates_storage.get_download_url.assert_awaited_once_with(template_id)

        response = await client.get(f"/api/application-templates/{not_uploaded_template_id}/download-url")

        assert response.status_code == status.HTTP_409_CONFLICT
    finally:
        await delete_templates(session_maker, template_id, not_uploaded_template_id)


async def test_update_application_template(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS)
    new_fields = [
        {**FIELDS_JSON[0], "label": "Wydział", "fieldType": "SELECT", "options": ["W4", "W8"], "defaultValue": "W4"},
        {**FIELDS_JSON[1], "label": "Numer albumu", "pattern": r"\d{6}", "passToAi": False},
    ]

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Nowa nazwa", "instructions": "Instrukcje", "fields": new_fields},
        )

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == {
            "id": str(template_id),
            "createDate": ANY,
            "name": "Nowa nazwa",
            "status": "DRAFT",
            "fields": new_fields,
            "instructions": "Instrukcje",
            "description": None,
            "descriptionStatus": "NOT_REQUESTED",
            "organizationId": None,
            "suborganizationId": None,
        }

        async with session_maker() as session:
            fields = await session.scalar(
                select(application_templates_table.c.fields).where(application_templates_table.c.id == template_id)
            )
        assert fields[0].field_type == ApplicationFieldType.SELECT
        assert fields[1].pass_to_ai is False
    finally:
        await delete_templates(session_maker, template_id)


async def test_update_application_template_with_invalid_fields(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Wniosek", "fields": [FIELDS_JSON[0]]},
        )

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    finally:
        await delete_templates(session_maker, template_id)


async def test_publish_and_unpublish_application_template(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS, instructions="Instrukcje")

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/publish")

        assert response.status_code == status.HTTP_200_OK
        assert response.json()["status"] == "PUBLISHED"

        response = await client.post(f"/api/application-templates/{template_id}/unpublish")

        assert response.status_code == status.HTTP_200_OK
        assert response.json()["status"] == "DRAFT"
    finally:
        await delete_templates(session_maker, template_id)


async def test_publish_application_template_without_instructions(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/publish")

        assert response.status_code == status.HTTP_409_CONFLICT
    finally:
        await delete_templates(session_maker, template_id)


async def test_delete_application_template(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_templates_storage,
    mock_application_templates_storage,
):
    template_id = await insert_template(session_maker)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    response = await client.delete(f"/api/application-templates/{template_id}")

    assert response.status_code == status.HTTP_204_NO_CONTENT
    mock_application_templates_storage.delete_template.assert_awaited_once_with(template_id)

    response = await client.delete(f"/api/application-templates/{template_id}")

    assert response.status_code == status.HTTP_404_NOT_FOUND


async def test_get_published_application_templates_filtered_by_organization(
    client, override_session_maker, session_maker, override_authorize_normal_user
):
    organization_id = await insert_organization(session_maker, "UWr")
    other_organization_id = await insert_organization(session_maker, "UMed")
    suborganization_id = await insert_suborganization(session_maker, organization_id, "Wydział Prawa")
    other_suborganization_id = await insert_suborganization(session_maker, organization_id, "Wydział Chemii")

    published = ApplicationTemplateStatus.PUBLISHED
    general_id = await insert_template(session_maker, "General", published)
    await insert_template(session_maker, "Organization", published, organization_id=organization_id)
    await insert_template(
        session_maker,
        "Law",
        published,
        organization_id=organization_id,
        suborganization_id=suborganization_id,
    )
    await insert_template(
        session_maker,
        "Chemistry",
        published,
        organization_id=organization_id,
        suborganization_id=other_suborganization_id,
    )
    await insert_template(session_maker, "Other organization", published, organization_id=other_organization_id)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get("/api/application-templates/published")

        assert response.status_code == status.HTTP_200_OK
        assert {template["name"] for template in response.json()} >= {
            "General",
            "Organization",
            "Law",
            "Chemistry",
            "Other organization",
        }

        response = await client.get(
            "/api/application-templates/published", params={"organizationId": str(organization_id)}
        )

        assert response.status_code == status.HTTP_200_OK
        assert {template["name"] for template in response.json()} == {"General", "Organization", "Law", "Chemistry"}

        response = await client.get(
            "/api/application-templates/published",
            params={"organizationId": str(organization_id), "suborganizationId": str(suborganization_id)},
        )

        assert response.status_code == status.HTTP_200_OK
        assert {template["name"] for template in response.json()} == {"General", "Organization", "Law"}
    finally:
        async with session_maker.begin() as session:
            await session.execute(
                delete(application_templates_table).where(
                    application_templates_table.c.organization_id.in_([organization_id, other_organization_id])
                )
            )
        await delete_templates(session_maker, general_id)
        await delete_organizations(session_maker, organization_id, other_organization_id)


async def test_get_published_application_templates_suborganization_without_organization(
    client, override_session_maker, override_authorize_normal_user
):
    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    response = await client.get("/api/application-templates/published", params={"suborganizationId": str(uuid4())})

    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT


async def test_update_application_template_with_organization(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    organization_id = await insert_organization(session_maker, "UWr")
    suborganization_id = await insert_suborganization(session_maker, organization_id, "Wydział Prawa")
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={
                "name": "Wniosek",
                "fields": FIELDS_JSON,
                "organizationId": str(organization_id),
                "suborganizationId": str(suborganization_id),
            },
        )

        assert response.status_code == status.HTTP_200_OK
        assert response.json()["organizationId"] == str(organization_id)
        assert response.json()["suborganizationId"] == str(suborganization_id)
    finally:
        await delete_templates(session_maker, template_id)
        await delete_organizations(session_maker, organization_id)


async def test_update_application_template_with_not_existing_organization(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Wniosek", "fields": FIELDS_JSON, "organizationId": str(uuid4())},
        )

        assert response.status_code == status.HTTP_404_NOT_FOUND
    finally:
        await delete_templates(session_maker, template_id)


async def test_update_application_template_with_suborganization_of_other_organization(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    organization_id = await insert_organization(session_maker, "UWr")
    other_organization_id = await insert_organization(session_maker, "UMed")
    other_suborganization_id = await insert_suborganization(session_maker, other_organization_id, "Wydział Lekarski")
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={
                "name": "Wniosek",
                "fields": FIELDS_JSON,
                "organizationId": str(organization_id),
                "suborganizationId": str(other_suborganization_id),
            },
        )

        assert response.status_code == status.HTTP_404_NOT_FOUND
    finally:
        await delete_templates(session_maker, template_id)
        await delete_organizations(session_maker, organization_id, other_organization_id)


async def test_delete_suborganization_keeps_application_template_in_organization(session_maker):
    organization_id = await insert_organization(session_maker, "UWr")
    suborganization_id = await insert_suborganization(session_maker, organization_id, "Wydział Prawa")
    template_id = await insert_template(
        session_maker, organization_id=organization_id, suborganization_id=suborganization_id
    )

    try:
        async with session_maker.begin() as session:
            await session.execute(
                delete(suborganizations_table).where(suborganizations_table.c.id == suborganization_id)
            )

        assert await get_template_organization(session_maker, template_id) == (organization_id, None)
    finally:
        await delete_templates(session_maker, template_id)
        await delete_organizations(session_maker, organization_id)


async def test_delete_organization_makes_application_template_general(session_maker):
    organization_id = await insert_organization(session_maker, "UWr")
    suborganization_id = await insert_suborganization(session_maker, organization_id, "Wydział Prawa")
    template_id = await insert_template(
        session_maker, organization_id=organization_id, suborganization_id=suborganization_id
    )

    try:
        await delete_organizations(session_maker, organization_id)

        assert await get_template_organization(session_maker, template_id) == (None, None)
    finally:
        await delete_templates(session_maker, template_id)


async def test_update_application_template_with_suborganization_without_organization(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(session_maker, fields=FIELDS)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Wniosek", "fields": FIELDS_JSON, "suborganizationId": str(uuid4())},
        )

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    finally:
        await delete_templates(session_maker, template_id)


async def test_search_published_application_templates(
    client, override_session_maker, session_maker, override_authorize_normal_user
):
    prawobiorca.dependency_overrides[get_texts_embedder] = lambda: StubTextsEmbedder()
    organization_id = await insert_organization(session_maker, "PWR")
    other_organization_id = await insert_organization(session_maker, "UWR")
    matching_id = await insert_template(
        session_maker,
        "Matching",
        ApplicationTemplateStatus.PUBLISHED,
        description="Wniosek o urlop.",
        description_vector=QUERY_VECTOR,
    )
    half_matching_id = await insert_template(
        session_maker,
        "Half matching",
        ApplicationTemplateStatus.PUBLISHED,
        organization_id=organization_id,
        description_vector=HALF_RELATED_VECTOR,
    )
    templates_ids = [
        matching_id,
        half_matching_id,
        await insert_template(
            session_maker, "Unrelated", ApplicationTemplateStatus.PUBLISHED, description_vector=UNRELATED_VECTOR
        ),
        await insert_template(session_maker, "Without description", ApplicationTemplateStatus.PUBLISHED),
        await insert_template(session_maker, "Draft", ApplicationTemplateStatus.DRAFT, description_vector=QUERY_VECTOR),
        await insert_template(
            session_maker,
            "Other organization",
            ApplicationTemplateStatus.PUBLISHED,
            organization_id=other_organization_id,
            description_vector=QUERY_VECTOR,
        ),
    ]

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get(
            "/api/application-templates/published/search",
            params={"query": "urlop", "threshold": 0.5, "organizationId": str(organization_id)},
        )

        assert response.status_code == status.HTTP_200_OK
        assert response.json() == [
            {
                "template": {
                    "id": str(matching_id),
                    "name": "Matching",
                    "description": "Wniosek o urlop.",
                    "fields": [],
                },
                "score": pytest.approx(1.0),
            },
            {
                "template": {"id": str(half_matching_id), "name": "Half matching", "description": None, "fields": []},
                "score": pytest.approx(1 / math.sqrt(2)),
            },
        ]

        response = await client.get(
            "/api/application-templates/published/search", params={"query": "urlop", "threshold": 0.9}
        )

        assert response.status_code == status.HTTP_200_OK
        assert [result["template"]["name"] for result in response.json()] == ["Matching", "Other organization"]
    finally:
        prawobiorca.dependency_overrides.pop(get_texts_embedder, None)
        await delete_templates(session_maker, *templates_ids)
        await delete_organizations(session_maker, organization_id, other_organization_id)


async def test_search_published_application_templates_empty_query(
    client, override_session_maker, override_authorize_normal_user
):
    prawobiorca.dependency_overrides[get_texts_embedder] = lambda: StubTextsEmbedder()

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.get(
            "/api/application-templates/published/search", params={"query": "", "threshold": 0.5}
        )

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    finally:
        prawobiorca.dependency_overrides.pop(get_texts_embedder, None)


async def test_regenerate_application_template_description(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_template_description_generation_scheduler,
    mock_application_template_description_generation_scheduler,
):
    template_id = await insert_template(session_maker, fields=FIELDS, instructions="Instrukcje")
    scheduler = mock_application_template_description_generation_scheduler

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/description-generation")

        assert response.status_code == status.HTTP_202_ACCEPTED
        scheduler.schedule_application_template_description_generation.assert_awaited_once_with(ANY, template_id)
        description = await get_template_description(session_maker, template_id)
        assert description.description_status == ApplicationTemplateDescriptionStatus.IN_PROGRESS

        response = await client.post(f"/api/application-templates/{template_id}/description-generation")

        assert response.status_code == status.HTTP_409_CONFLICT
    finally:
        await delete_templates(session_maker, template_id)

    response = await client.post(f"/api/application-templates/{template_id}/description-generation")

    assert response.status_code == status.HTTP_404_NOT_FOUND


async def test_regenerate_published_application_template_description(
    client,
    override_session_maker,
    session_maker,
    override_authorize_admin_user,
    override_get_application_template_description_generation_scheduler,
):
    template_id = await insert_template(
        session_maker, template_status=ApplicationTemplateStatus.PUBLISHED, instructions="Instrukcje"
    )

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/description-generation")

        assert response.status_code == status.HTTP_409_CONFLICT
    finally:
        await delete_templates(session_maker, template_id)


async def test_regenerate_application_template_description_as_normal_user(
    client,
    override_session_maker,
    session_maker,
    override_authorize_normal_user,
    override_get_application_template_description_generation_scheduler,
):
    template_id = await insert_template(session_maker)

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.post(f"/api/application-templates/{template_id}/description-generation")

        assert response.status_code == status.HTTP_403_FORBIDDEN
    finally:
        await delete_templates(session_maker, template_id)


async def test_update_application_template_description_clears_description_vector(
    client, override_session_maker, session_maker, override_authorize_admin_user
):
    template_id = await insert_template(
        session_maker, fields=FIELDS, description="Wniosek o urlop.", description_vector=QUERY_VECTOR
    )

    client.cookies.set(ACCESS_COOKIE_NAME, ACCESS_TOKEN)

    try:
        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Wniosek", "description": "Wniosek o urlop.", "fields": FIELDS_JSON},
        )

        assert response.status_code == status.HTTP_200_OK
        description = await get_template_description(session_maker, template_id)
        assert description.description_vector is not None

        response = await client.put(
            f"/api/application-templates/{template_id}",
            json={"name": "Wniosek", "description": "Wniosek o stypendium.", "fields": FIELDS_JSON},
        )

        assert response.status_code == status.HTTP_200_OK
        assert response.json()["description"] == "Wniosek o stypendium."
        description = await get_template_description(session_maker, template_id)
        assert description.description_vector is None
    finally:
        await delete_templates(session_maker, template_id)
